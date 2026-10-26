"""Excel generation service using openpyxl.

Produces styled, professional Microsoft Excel (.xlsx) workbooks for
faculty attendance reports and assignment submissions with conditional formatting,
auto-fitted columns, and calculation summaries.
"""

from datetime import datetime
import io
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

BANNER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
BANNER_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")

SUB_FONT = Font(name="Calibri", size=9, italic=True, color="64748B")

SAFE_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
SAFE_FONT = Font(name="Calibri", size=10, bold=True, color="166534")

WARN_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WARN_FONT = Font(name="Calibri", size=10, bold=True, color="B45309")

CRIT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
CRIT_FONT = Font(name="Calibri", size=10, bold=True, color="991B1B")

REG_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

TOTALS_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")


def build_attendance_xlsx(subject_dict: Dict[str, Any], records: List[Dict[str, Any]]) -> bytes:
    """Generate a styled Excel workbook for subject attendance."""
    wb = openpyxl.Workbook()
    ws = wb.active
    code = subject_dict.get("code", "ATTENDANCE")
    ws.title = f"{code[:25]} Attendance"

    # Row 1: Brand / Title banner
    ws.merge_cells("A1:N1")
    b_cell = ws["A1"]
    b_cell.value = f"CampusPilot AI — Attendance Report: {subject_dict.get('name', '')} ({code})"
    b_cell.font = BANNER_FONT
    b_cell.fill = BANNER_FILL
    b_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # Row 2: Subtitle
    ws.merge_cells("A2:N2")
    sub_cell = ws["A2"]
    dept = subject_dict.get("department", "CSE")
    sem = subject_dict.get("semester", 1)
    now_str = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    sub_cell.value = f"Department: {dept}   |   Semester: {sem}   |   Total Enrolled: {len(records)}   |   Exported: {now_str}"
    sub_cell.font = SUB_FONT
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Row 3: Blank separator
    ws.append([])
    ws.row_dimensions[3].height = 8

    # Row 4: Column Headers
    headers = [
        "Roll No",
        "Student Name",
        "Email",
        "Department",
        "Semester",
        "Section",
        "Subject Code",
        "Subject Name",
        "Total Classes",
        "Attended",
        "Missed",
        "Attendance %",
        "Status",
        "Eligibility Guidance",
    ]
    ws.append(headers)
    ws.row_dimensions[4].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # Rows 5+: Student Records
    start_row = 5
    for idx, r in enumerate(records, start=start_row):
        pct = float(r.get("percentage", 0.0))
        status = str(r.get("status", "NO_RECORD")).upper()

        row_values = [
            r.get("student_id", ""),
            r.get("full_name", ""),
            r.get("email", ""),
            r.get("department", ""),
            r.get("semester", ""),
            r.get("section", ""),
            subject_dict.get("code", ""),
            subject_dict.get("name", ""),
            int(r.get("total_classes", 0)),
            int(r.get("attended_classes", 0)),
            int(r.get("total_classes", 0)) - int(r.get("attended_classes", 0)),
            f"{pct:.2f}%",
            status,
            r.get("guidance", ""),
        ]
        ws.append(row_values)
        ws.row_dimensions[idx].height = 20

        # Style data cells
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.font = REG_FONT
            c.border = THIN_BORDER
            c.alignment = Alignment(
                horizontal="right" if col_idx in (1, 5, 9, 10, 11, 12) else "left",
                vertical="center",
            )

        # Highlight Status column
        stat_cell = ws.cell(row=idx, column=13)
        stat_cell.alignment = Alignment(horizontal="center", vertical="center")
        if "SAFE" in status:
            stat_cell.fill = SAFE_FILL
            stat_cell.font = SAFE_FONT
        elif "WARN" in status:
            stat_cell.fill = WARN_FILL
            stat_cell.font = WARN_FONT
        elif "CRIT" in status:
            stat_cell.fill = CRIT_FILL
            stat_cell.font = CRIT_FONT

    # Summary Row
    end_row = start_row + len(records)
    ws.row_dimensions[end_row].height = 22
    summary_label = ws.cell(row=end_row, column=2, value="AVERAGE / TOTALS")
    summary_label.font = BOLD_FONT
    summary_label.alignment = Alignment(horizontal="left", vertical="center")

    avg_pct = (
        sum(float(r.get("percentage", 0)) for r in records) / len(records)
        if records
        else 0.0
    )
    avg_cell = ws.cell(row=end_row, column=12, value=f"{avg_pct:.2f}%")
    avg_cell.font = BOLD_FONT
    avg_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=end_row, column=col_idx)
        cell.fill = TOTALS_FILL
        cell.border = THIN_BORDER

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title and subtitle merged rows
            if cell.row in (1, 2, 3):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_submissions_xlsx(assignment_dict: Dict[str, Any], submissions: List[Dict[str, Any]]) -> bytes:
    """Generate a styled Excel workbook for assignment submissions."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Submissions"

    # Row 1: Title
    ws.merge_cells("A1:K1")
    b_cell = ws["A1"]
    b_cell.value = f"CampusPilot AI — Submissions: {assignment_dict.get('title', '')}"
    b_cell.font = BANNER_FONT
    b_cell.fill = BANNER_FILL
    b_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # Row 2: Subtitle
    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    due = assignment_dict.get("due_date", "")
    code = assignment_dict.get("subject_code", "")
    sub_cell.value = f"Course: {code}   |   Due Date: {due}   |   Total Submissions: {len(submissions)}   |   Exported: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    sub_cell.font = SUB_FONT
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Row 3: Blank separator
    ws.append([])
    ws.row_dimensions[3].height = 8

    # Row 4: Headers
    headers = [
        "Roll No",
        "Student Name",
        "Email",
        "Department",
        "Semester",
        "Section",
        "Assignment Title",
        "Due Date",
        "Status",
        "Submission Timestamp",
        "Attachment / File URL",
    ]
    ws.append(headers)
    ws.row_dimensions[4].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # Rows 5+: Submissions
    for idx, s in enumerate(submissions, start=5):
        status = str(s.get("status", "submitted")).upper()
        row_values = [
            s.get("roll_no", ""),
            s.get("student_name", ""),
            s.get("email", ""),
            s.get("department", ""),
            s.get("semester", ""),
            s.get("section", ""),
            assignment_dict.get("title", ""),
            str(assignment_dict.get("due_date", "")),
            status,
            s.get("submitted_at", ""),
            s.get("file_url", "N/A"),
        ]
        ws.append(row_values)
        ws.row_dimensions[idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.font = REG_FONT
            c.border = THIN_BORDER
            c.alignment = Alignment(
                horizontal="right" if col_idx in (1, 5) else "left",
                vertical="center",
            )

        # Highlight status
        stat_cell = ws.cell(row=idx, column=9)
        stat_cell.alignment = Alignment(horizontal="center", vertical="center")
        if "APPROV" in status or "ACCEPT" in status:
            stat_cell.fill = SAFE_FILL
            stat_cell.font = SAFE_FONT
        elif "REVIEW" in status:
            stat_cell.fill = WARN_FILL
            stat_cell.font = WARN_FONT

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
